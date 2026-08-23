from __future__ import annotations

import csv
import json
import sqlite3
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook


def _columns(rows: list[dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key not in seen:
                seen.add(key)
                columns.append(key)
    return columns


def _scalar(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def _spreadsheet_scalar(value: Any) -> Any:
    """Keep untrusted text from becoming a formula in CSV/Excel clients."""
    scalar = _scalar(value)
    if isinstance(scalar, str) and scalar.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + scalar
    return scalar


def write_exports(run_dir: Path, rows: list[dict[str, Any]]) -> dict[str, Path]:
    run_dir.mkdir(parents=True, exist_ok=True)
    columns = _columns(rows)
    paths = {
        "json": run_dir / "data.json",
        "jsonl": run_dir / "data.jsonl",
        "csv": run_dir / "data.csv",
        "xlsx": run_dir / "data.xlsx",
        "sqlite": run_dir / "data.sqlite",
    }
    paths["json"].write_text(json.dumps(rows, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    with paths["jsonl"].open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")
    with paths["csv"].open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows({key: _spreadsheet_scalar(value) for key, value in row.items()} for row in rows)

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("data")
    sheet.append(columns)
    for row in rows:
        sheet.append([_spreadsheet_scalar(row.get(column)) for column in columns])
    workbook.save(paths["xlsx"])

    with sqlite3.connect(paths["sqlite"]) as connection:
        connection.execute("DROP TABLE IF EXISTS data")
        if columns:
            quoted = ", ".join(f'"{column.replace(chr(34), chr(34) * 2)}" TEXT' for column in columns)
            connection.execute(f"CREATE TABLE data ({quoted})")
            placeholders = ",".join("?" for _ in columns)
            connection.executemany(
                f"INSERT INTO data VALUES ({placeholders})",
                [[None if row.get(column) is None else str(_scalar(row.get(column))) for column in columns] for row in rows],
            )
            connection.commit()
    return paths


def write_bundle(run_dir: Path) -> Path:
    bundle = run_dir / "weaver-bundle.zip"
    with zipfile.ZipFile(bundle, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(run_dir.rglob("*")):
            if path.is_file() and path != bundle:
                archive.write(path, path.relative_to(run_dir))
    return bundle
